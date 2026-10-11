import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from docx_helpers import (
    set_cell_background,
    set_cell_margins,
    set_table_borders,
    format_row,
    add_academic_heading,
    add_body_p,
    add_bullet_p
)

IMG_DIR = os.path.abspath("report_images")

def add_figure_image(doc, img_name, caption_text, width_inches=5.6):
    img_path = os.path.join(IMG_DIR, img_name)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        rc = p_cap.add_run(caption_text)
        rc.bold = True
        rc.font.name = 'Times New Roman'
        rc.font.size = Pt(11)
        rc.font.color.rgb = RGBColor(0, 0, 0)
    else:
        print(f"Warning: Image {img_path} not found!")

def add_algorithm_step(doc, num_str, step_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25)
    
    rn = p.add_run(num_str + " ")
    rn.bold = True
    rn.font.name = 'Times New Roman'
    rn.font.size = Pt(11.5)
    
    rt = p.add_run(step_text)
    rt.font.name = 'Times New Roman'
    rt.font.size = Pt(11.5)
    rt.font.color.rgb = RGBColor(0, 0, 0)

def build_perfect_report():
    doc = Document()

    # Exact standard margins: 1.25 in left (for thesis binding), 1.0 in top, bottom, right
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        
        # Centered page number in footer
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        p_ft._p.append(fldSimple)

    # Base typography: Times New Roman, 12pt, Pure Black
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)

    # =========================================================================
    # 1. TITLE PAGE (Page 1)
    # =========================================================================
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(36)
    p1.paragraph_format.space_after = Pt(12)
    r = p1.add_run("COGNIVISTA — ADAPTIVE BRAIN CHALLENGE &\nCOGNITIVE TRAINING AGENT USING NLP")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor(0, 0, 0)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(36)
    r = p2.add_run("23IT23E – AI TOOLS FOR NATURAL LANGUAGE PROCESSING")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(6)
    r = p3.add_run("Submitted by")
    r.italic = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_after = Pt(40)
    r = p4.add_run("[STUDENT NAME]          [REGISTER NUMBER]")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)

    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_after = Pt(6)
    r = p5.add_run("In partial fulfillment for the award of the degree of")
    r.italic = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p6 = doc.add_paragraph()
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p6.paragraph_format.space_after = Pt(4)
    r = p6.add_run("BACHELOR OF TECHNOLOGY")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)

    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p7.paragraph_format.space_after = Pt(70)
    r = p7.add_run("in\nINFORMATION TECHNOLOGY")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)

    p8 = doc.add_paragraph()
    p8.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p8.paragraph_format.space_after = Pt(4)
    r = p8.add_run("NATIONAL ENGINEERING COLLEGE\n(An Autonomous Institution affiliated to Anna University, Chennai)\nK.R.NAGAR, KOVILPATTI - 628503\n[OCTOBER 2026]")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)

    doc.add_page_break()

    # =========================================================================
    # 2. BONAFIDE CERTIFICATE (Page 2)
    # =========================================================================
    p_bc1 = doc.add_paragraph()
    p_bc1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bc1.paragraph_format.space_before = Pt(24)
    p_bc1.paragraph_format.space_after = Pt(4)
    r = p_bc1.add_run("NATIONAL ENGINEERING COLLEGE\n(An Autonomous Institution affiliated to Anna University, Chennai)\nK.R. NAGAR, KOVILPATTI - 628503")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)

    p_bc2 = doc.add_paragraph()
    p_bc2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bc2.paragraph_format.space_before = Pt(36)
    p_bc2.paragraph_format.space_after = Pt(36)
    r = p_bc2.add_run("BONAFIDE CERTIFICATE")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)

    add_body_p(
        doc,
        'This is to certify that this project report, "COGNIVISTA — Adaptive Brain Challenge & Cognitive Training Agent Using NLP", is the bonafide work of [STUDENT NAME] ([REGISTER NUMBER]) who carried out the project work under my supervision for the course 23IT23E – AI Tools for Natural Language Processing during the academic year 2026.',
        space_after=180
    )

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(60)
    p_sig.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_left = p_sig.add_run("Course Instructor / Guide                                                     Head of the Department")
    r_left.bold = True
    r_left.font.name = 'Times New Roman'
    r_left.font.size = Pt(12)

    doc.add_page_break()

    # =========================================================================
    # 3. TABLE OF CONTENTS (Page 3)
    # =========================================================================
    add_academic_heading(doc, "TABLE OF CONTENTS", level=1)
    
    p_toc_note = doc.add_paragraph()
    p_toc_note.paragraph_format.space_after = Pt(12)

    toc_table = doc.add_table(rows=11, cols=3)
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(toc_table)
    
    hdr_cells = toc_table.rows[0].cells
    hdr_cells[0].paragraphs[0].add_run("CHAPTER NO").bold = True
    hdr_cells[1].paragraphs[0].add_run("TITLE").bold = True
    hdr_cells[2].paragraphs[0].add_run("PAGE NO").bold = True
    format_row(toc_table.rows[0], is_header=True)
    
    toc_entries = [
        ("1", "INTRODUCTION", "4"),
        ("2", "OBJECTIVES", "6"),
        ("3", "SYSTEM DESCRIPTION & NLP APPROACH", "8"),
        ("4", "AI & NLP PIPELINE ARCHITECTURE", "11"),
        ("5", "MATHEMATICAL FORMULATIONS & ALGORITHMS", "15"),
        ("6", "TECH STACK & PACKAGES", "19"),
        ("7", "DATABASE & KNOWLEDGE BASE STRUCTURES", "22"),
        ("8", "MODULES & IMPLEMENTATION WALKTHROUGH", "27"),
        ("9", "OUTPUT SCREENSHOTS", "31"),
        ("10", "CONCLUSION & FUTURE ENHANCEMENTS", "35"),
    ]
    
    col_widths = [Inches(1.2), Inches(4.3), Inches(0.9)]
    for idx, (ch_no, ch_title, ch_page) in enumerate(toc_entries, start=1):
        row = toc_table.rows[idx]
        row.cells[0].paragraphs[0].add_run(ch_no)
        row.cells[1].paragraphs[0].add_run(ch_title)
        row.cells[2].paragraphs[0].add_run(ch_page)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = col_widths[i]

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 1: INTRODUCTION (Pages 4, 5, and ~3/4 of Page 6)
    # =========================================================================
    add_academic_heading(doc, "CHAPTER 1\nINTRODUCTION", level=1)
    
    add_body_p(
        doc,
        "Human cognitive faculties—ranging from working memory capacity and executive processing speed to spatial visualization, deductive logic, and lateral reasoning—are essential for academic performance, professional competence, and daily analytical decision-making. Over the past decade, digital cognitive assessment and brain-training platforms have gained widespread traction across mobile and web environments. Millions of individuals regularly engage with digital puzzle games with the goal of sharpening their mental acuity, enhancing problem-solving agility, and counteracting age-related cognitive decline."
    )
    add_body_p(
        doc,
        "However, the vast majority of existing brain-training software suffers from severe structural rigidity: challenges are drawn from static, hardcoded question banks, interactive visual representations are scarce or disconnected from dynamic problem variables, and user evaluation is confined to binary correct/incorrect tallies without explainable linguistic diagnosis. When a puzzle platform relies on a closed set of pre-written questions, users quickly experience item fatigue, shifting their problem-solving behavior from genuine cognitive synthesis to superficial rote recall."
    )
    add_body_p(
        doc,
        "Cognivista represents a paradigm shift in cognitive training systems by deploying an autonomous Artificial Intelligence and Natural Language Processing (NLP) agent at the center of an interactive multi-domain challenge universe. Rather than serving static questions, Cognivista dynamically synthesizes original multi-genre cognitive challenges on demand, pairs them with procedural vector visual schematics (such as geometric triangles with angle labels, electrical circuits with toggled switches, and rotating gear trains), and continuously analyzes user response times, error patterns, and cognitive bottlenecks. By synthesizing symbolic parsing rules with large language model reasoning, Cognivista delivers an intelligent, adaptive, and personalized training experience that evolves alongside the learner."
    )

    add_academic_heading(doc, "1.1 Problem Statement", level=2)
    add_body_p(
        doc,
        "Conventional computerized cognitive training applications exhibit critical limitations that hinder their educational and neurological efficacy:"
    )
    add_bullet_p(doc, " Most platforms rely on pre-compiled lists of trivia or formulaic logic puzzles. Users quickly exhaust the item pool, leading to rote memorization rather than authentic cognitive exercise.", bold_title="• Item Pool Exhaustion:")
    add_bullet_p(doc, " Challenges are predominantly rendered as plain sentence text or multiple-choice paragraphs. Real-world cognitive problem solving frequently demands visual-spatial and diagrammatic manipulation, which is absent in text-only systems.", bold_title="• Absence of Dynamic Visual Representation:")
    add_bullet_p(doc, " Standard applications simply award a point or display a red cross when an answer is wrong. They fail to explain the underlying cognitive failure mechanism (e.g., calculation slippage, misread premise, spatial inversion, or distractor entrapment).", bold_title="• Lack of Diagnostic Pedagogical Explanations:")
    add_bullet_p(doc, " Systems adjust difficulty through blunt linear increments (Level 1, Level 2, Level 3) without tracking granular, domain-specific competencies or cognitive fatigue curves.", bold_title="• Inelastic and Coarse Difficulty Progression:")
    add_bullet_p(doc, " Puzzles exist in an artificial vacuum isolated from real-world temporal, geographic, and economic realities, diminishing practical analytical transfer.", bold_title="• Isolation from Real-World Intelligence:")

    add_academic_heading(doc, "1.2 Proposed Solution", level=2)
    add_body_p(
        doc,
        "The proposed system, COGNIVISTA, resolves these limitations by uniting state-of-the-art Natural Language Processing, generative multi-agent orchestration, procedural vector rendering, and adaptive psychometric scoring into an end-to-end cognitive training desk:"
    )
    add_bullet_p(doc, " An autonomous agent parses user intent, cognitive weaknesses, and requested genres to dynamically generate endless unique challenges across 12 distinct cognitive domains.", bold_title="• Dynamic Challenge Universe:")
    add_bullet_p(doc, " Every visual puzzle is synthesized with procedural vector SVG schematics—including geometric angle networks, series-parallel circuits, interlocking gear trains, balanced scales, and flight trajectory maps—with mathematically consistent parameters.", bold_title="• Procedural Visual Schematics:")
    add_bullet_p(doc, " A hybrid symbolic and neural evaluator assesses responses, classifies errors into pedagogical failure taxonomies, and powers an interactive conversational desk where users can ask 'Why was my answer wrong?' to receive step-by-step Socratic remediation.", bold_title="• Diagnostic Response Evaluation & Interactive Review:")
    add_bullet_p(doc, " A continuous feedback loop computes domain-specific Elo proficiency ratings, tracks reaction latencies, and dynamically regulates challenge difficulty, pacing, and domain switching in real time.", bold_title="• Adaptive Psychometric Training Loop:")
    add_bullet_p(doc, " Live external APIs (Open-Meteo Weather, World Time API, Consumer Market indexes) are integrated with verified source attribution and freshness timestamps to produce timely, contextual analytical challenges.", bold_title="• Real-Time Environmental Grounding:")

    add_academic_heading(doc, "1.3 Scope of the Project", level=2)
    add_body_p(
        doc,
        "The scope of the Cognivista project encompasses the architectural design, algorithmic implementation, and empirical verification of a complete full-stack web application. The backend is engineered using Python 3.11, FastAPI, and SQLAlchemy, deploying Google Gemini 2.5 Flash through the official Google GenAI SDK alongside deterministic rule-based NLP filters. The frontend is built using React 19, TypeScript, Vite, and high-performance vector SVG rendering. The application features six functionally dedicated operational modules: Cognitive Progress Dashboard, Natural-Language Challenge Creator Desk, Interactive Challenge Arena, Post-Attempt Review & Diagnostic Desk, Adaptive Training Session Loop, and Live Grounded Intelligence Feeds. Scanned physical documents, legacy proprietary file formats, and clinical psychiatric diagnoses lie outside the current engineering scope."
    )

    add_academic_heading(doc, "1.4 Key Architectural Contributions", level=2)
    add_body_p(
        doc,
        "To accomplish the goals of an autonomous cognitive challenge platform, Cognivista introduces several key technical and pedagogical innovations that distinguish it from commercial quiz engines and static web portals:"
    )
    add_bullet_p(
        doc,
        " Unlike traditional platforms that embed static image files that cannot be altered dynamically, Cognivista computes vector coordinates in real time. Geometric vertices, angle arcs, switch terminals, wire paths, and gear teeth are generated on the fly as mathematically consistent SVG objects.",
        bold_title="1. Dynamic Parametric Vector Synthesis: "
    )
    add_bullet_p(
        doc,
        " The system decouples natural language intent parsing into high-speed regex token extraction and neural few-shot slot resolution. This hybrid design achieves sub-millisecond execution for standard keyword commands while maintaining full semantic flexibility for complex, freeform user requests.",
        bold_title="2. Hybrid Symbolic-Neural Ingestion Pipeline: "
    )
    add_bullet_p(
        doc,
        " When a user errs, Cognivista avoids blunt negative reinforcement. Instead, it compares the user's answer against pre-computed distractor models, classifies the underlying misconception (such as arithmetic drift or premise negation), and initiates a guided Socratic dialogue.",
        bold_title="3. Diagnostic Cognitive Error Taxonomy: "
    )
    add_bullet_p(
        doc,
        " The cognitive agent operates with an open, inspectable decision stream. Before dispatching any challenge during an adaptive workout, it records its pedagogical rationale, user fatigue score, and targeted skill deficit, fostering user metacognition and transparency.",
        bold_title="4. Inspectable Autonomous Agent Telemetry: "
    )

    add_academic_heading(doc, "1.5 Organization of the Report", level=2)
    add_body_p(
        doc,
        "The remainder of this report is structured into nine sequential chapters: Chapter 2 formalizes the specific engineering and pedagogical objectives. Chapter 3 presents the system description, modular breakdown, and hybrid NLP methodology. Chapter 4 delineates the end-to-end pipeline and layered architecture. Chapter 5 articulates the mathematical formulations governing psychometric Elo scoring, latency ratios, and the core algorithmic step procedures. Chapter 6 outlines the software tech stack and runtime dependencies. Chapter 7 details the relational database schema and knowledge bases. Chapter 8 provides an implementation walkthrough with source code. Chapter 9 presents the application output screenshots. Finally, Chapter 10 concludes the report and identifies avenues for future enhancement."
    )

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 2: OBJECTIVES
    # =========================================================================
    add_academic_heading(doc, "CHAPTER 2\nOBJECTIVES", level=1)
    
    add_body_p(
        doc,
        "The Cognivista project was conceived and implemented to address the technological and pedagogical limitations identified in contemporary computer-assisted cognitive training systems. To establish a rigorous engineering benchmark, the objectives are categorized into three core domains: Engineering & Algorithmic, Cognitive & Pedagogical, and User Experience & Metacognitive."
    )

    add_academic_heading(doc, "2.1 Primary Engineering & Algorithmic Objectives", level=2)
    add_bullet_p(doc, "To design and deploy an asynchronous, high-throughput backend server using Python 3.11 and FastAPI capable of serving challenge generation and evaluation requests within a sub-second response threshold.", bold_title="1. High-Throughput RESTful Architecture: ")
    add_bullet_p(doc, "To engineer a dynamic multi-genre challenge synthesis engine that generates original, mathematically verified problems across 12 distinct intellectual faculties without drawing from static question pools.", bold_title="2. Autonomous Procedural Problem Synthesis: ")
    add_bullet_p(doc, "To implement a procedural SVG vector rendering pipeline that converts abstract mathematical parameters into interactive geometric angle diagrams, circuit schematics, interlocking gear trains, and balance scales.", bold_title="3. Procedural Vector Schematic Generation: ")
    add_bullet_p(doc, "To construct a hybrid NLP intent parser combining compiled regular expressions with few-shot Large Language Model (Gemini 2.5 Flash) slot-filling to reliably extract domain, difficulty, and format parameters.", bold_title="4. Robust Intent Parsing & Slot Filling: ")
    add_bullet_p(doc, "To formulate and implement a multi-dimensional Elo rating algorithm that maintains independent psychometric competence scores across seven distinct cognitive domains with dynamic sensitivity scaling.", bold_title="5. Multi-Domain Psychometric Rating Engine: ")

    add_academic_heading(doc, "2.2 Cognitive & Pedagogical Objectives", level=2)
    add_bullet_p(doc, "To formulate a comprehensive cognitive error taxonomy that categorizes user errors into five distinct failure mechanisms: Calculation Slippage, Misread Premise, Logical Fallacy, Spatial Inversion, and Distractor Trap.", bold_title="6. Diagnostic Error Taxonomy Classification: ")
    add_bullet_p(doc, "To develop an interactive conversational Socratic review module ('Why was my answer wrong?') that provides targeted conceptual hints and guided remediation without prematurely revealing direct solutions.", bold_title="7. Socratic Explainability & Remediation: ")
    add_bullet_p(doc, "To construct an adaptive session controller that modulates challenge difficulty, regulates pacing, and prevents cognitive fatigue by switching genres based on moving-average performance metrics.", bold_title="8. Dynamic Working Memory & Scaffolding: ")
    add_bullet_p(doc, "To ground abstract analytical reasoning in real-world physical and economic contexts by fetching live weather, time-zone, and market metrics with verified attribution metadata.", bold_title="9. Real-World Live Environmental Grounding: ")

    add_academic_heading(doc, "2.3 User Experience & Metacognitive Objectives", level=2)
    add_bullet_p(doc, "To engineer an inspectable, real-time AI Agent Autonomous Decision Stream log that exposes the agent's internal reasoning, target skill deficits, and difficulty modulations to the learner.", bold_title="10. Transparent Autonomous Decision Stream: ")
    add_bullet_p(doc, "To implement seven distinct interactive problem-solving arenas (matching pairs, sequence ordering, multiple selection, true/false switches, fill-in-the-blank, odd-one-out, and detective clues).", bold_title="11. Versatile Interactive Arena Paradigms: ")
    add_bullet_p(doc, "To provide an analytics dashboard that tracks historical attempt accuracy, response time velocity distributions, cognitive domain strengths, and streak progress.", bold_title="12. Longitudinal Cognitive Telemetry: ")

    add_academic_heading(doc, "2.4 Measurable Success Criteria & Verification Targets", level=2)
    add_body_p(
        doc,
        "To ensure verifiable implementation rigor, Table 2.1 defines the quantitative engineering benchmarks established for the Cognivista platform."
    )

    # Table 2.1
    p_t21 = doc.add_paragraph()
    p_t21.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t21.paragraph_format.space_before = Pt(6)
    p_t21.paragraph_format.space_after = Pt(4)
    r = p_t21.add_run("Table 2.1: Quantitative Engineering & Pedagogical Benchmarks")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t21 = doc.add_table(rows=6, cols=3)
    t21.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t21)
    
    t21_headers = ["Engineering Parameter", "Target Metric", "Verification Methodology"]
    for i, h in enumerate(t21_headers):
        t21.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t21.rows[0], is_header=True)
    
    t21_data = [
        ("Intent Parsing Latency", "< 50 ms for regex, < 650 ms for LLM", "Asynchronous timing benchmarks across 100 sample prompts"),
        ("Challenge Generation Variety", "12 distinct cognitive genres supported", "Automated test suite asserting valid schema generation per genre"),
        ("Vector Schematic Render Time", "< 15 ms client-side SVG rendering", "Chrome Performance Profiler measuring DOM paint events"),
        ("Diagnostic Error Classification", "> 92% accurate categorization of mistakes", "Evaluation against synthetically constructed distractor error sets"),
        ("Client Build Verification", "Zero TypeScript compilation warnings", "Automated CI build execution (`tsc -b && vite build`)")
    ]
    
    t21_widths = [Inches(2.2), Inches(2.3), Inches(2.7)]
    for r_idx, row_data in enumerate(t21_data, start=1):
        row = t21.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t21_widths[i]

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 3: SYSTEM DESCRIPTION & NLP APPROACH
    # =========================================================================
    add_academic_heading(doc, "CHAPTER 3\nSYSTEM DESCRIPTION & NLP APPROACH", level=1)
    
    add_academic_heading(doc, "3.1 System Overview", level=2)
    add_body_p(
        doc,
        "Cognivista is an enterprise-grade, full-stack cognitive computing and brain-training application designed to enhance human analytical, logical, and spatial reasoning skills through autonomous artificial intelligence. Operating as an intelligent pedagogical agent, Cognivista continuously observes the learner's performance across diverse intellectual challenges, diagnoses conceptual weaknesses, generates custom training scenarios, and delivers actionable linguistic feedback. The application is accessed via a responsive web client and powered by a high-throughput asynchronous backend server."
    )

    add_academic_heading(doc, "3.2 System Modules", level=2)
    add_body_p(
        doc,
        "The Cognivista application architecture is partitioned into six dedicated functional modules, each engineered for a specialized operational role:"
    )

    # Distinct module presentations with clear formatting
    add_academic_heading(doc, "Module 1: Cognitive Progress & Mastery Analytics Desk", level=3)
    add_body_p(
        doc,
        "This module serves as the learner's cognitive command center. It continuously aggregates historical attempt telemetry, computes domain-specific Elo ratings across seven core faculties, plots response latency distributions, highlights cognitive bottlenecks, and presents an interactive workout activity calendar tracking daily intellectual exertion."
    )

    add_academic_heading(doc, "Module 2: Autonomous Challenge Creator & Prompt Engineering Desk", level=3)
    add_body_p(
        doc,
        "This module empowers learners to submit freeform natural-language instructions (e.g., 'Generate an electric circuit challenge with parallel branches and switches'). The agent executes a 4-step reasoning process (Intent Analysis -> Domain Routing -> Parameter Synthesis -> Vector Payload Construction) and provides an instant preview before practice."
    )

    add_academic_heading(doc, "Module 3: Dynamic Multi-Format Challenge Arena & Vector Engine", level=3)
    add_body_p(
        doc,
        "This active problem-solving environment supports 12 challenge genres and seven distinct interaction paradigms (matching, ordering, multi-select, true/false, fill-in-the-blank, odd-one-out, and detective deduction). It integrates procedural vector SVG graphics for visual problems."
    )

    add_academic_heading(doc, "Module 4: Post-Attempt Review & Diagnostic Socratic Desk", level=3)
    add_body_p(
        doc,
        "This module delivers in-depth post-attempt analysis, displaying side-by-side comparisons of submitted vs. correct answers, latency metrics, and error classifications. It features an interactive 'Why was my answer wrong?' conversational module for Socratic remediation."
    )

    add_academic_heading(doc, "Module 5: Autonomous Adaptive Training Workout Session Loop", level=3)
    add_body_p(
        doc,
        "An automated training regime where the AI agent manages multi-step cognitive workouts. The system displays a live AI AGENT AUTONOMOUS DECISION STREAM log detailing why specific difficulties and genres were chosen."
    )

    add_academic_heading(doc, "Module 6: Live Environmental Intelligence & Grounded Data Engine", level=3)
    add_body_p(
        doc,
        "This module ingests real-world time-series data from weather stations, world time zones, and economic market indices to construct timely, physically grounded mathematical and logical puzzles with verified attribution."
    )

    add_academic_heading(doc, "3.3 Hybrid NLP Approach", level=2)
    add_body_p(
        doc,
        "Cognivista employs a hybrid NLP methodology that synthesizes deterministic, rule-based symbolic parsing with probabilistic neural language model reasoning. The techniques used are summarized in Table 3.1."
    )

    # Table 3.1
    p_t31 = doc.add_paragraph()
    p_t31.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t31.paragraph_format.space_before = Pt(8)
    p_t31.paragraph_format.space_after = Pt(4)
    r = p_t31.add_run("Table 3.1: Natural Language Processing Techniques Employed in Cognivista")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t31 = doc.add_table(rows=7, cols=3)
    t31.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t31)
    
    t31_headers = ["NLP Technique", "Operational Role in Cognivista", "Underlying Implementation"]
    for i, h in enumerate(t31_headers):
        t31.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t31.rows[0], is_header=True)
    
    t31_data = [
        ("Tokenisation & Regex Filtering", "Extracts user commands, genres, and difficulty levels from input text", "Whitespace & punctuation tokenizers, compiled regex pattern dictionaries"),
        ("Intent Slot Extraction", "Maps natural prompts to target domain, complexity, and format variables", "Rule-based grammar matchers with Gemini Flash fallback for ambiguous syntax"),
        ("Few-Shot Prompt Engineering", "Guides generative LLM to synthesize valid JSON challenge payloads", "Pydantic structured output schemas with validated system instruction boundaries"),
        ("Fuzzy Lexical Normalization", "Evaluates open-ended user answers against acceptable semantic variants", "Levenshtein distance, case-folding, whitespace trimming, and synonym maps"),
        ("Semantic Error Classification", "Diagnoses underlying cognitive failure modes from incorrect attempts", "Rule-based discrepancy taxonomy cross-referenced with domain distractor models"),
        ("Attribution Metadata Parsing", "Extracts timestamps, coordinates, and publisher credits from live APIs", "JSON extraction, ISO-8601 string parsing, and citation formatting filters")
    ]
    
    t31_widths = [Inches(1.8), Inches(2.4), Inches(2.2)]
    for r_idx, row_data in enumerate(t31_data, start=1):
        row = t31.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t31_widths[i]

    add_academic_heading(doc, "3.4 Autonomous Agentic Decision-Making Principle", level=2)
    add_body_p(
        doc,
        "At the core of Cognivista's architecture is the principle of Explainable Autonomous Pedagogy. Unlike black-box neural networks that make opaque recommendations, Cognivista logs every pedagogical decision to an inspectable telemetry stream. Before selecting a challenge, the agent articulates: (1) current user rating in the target domain, (2) recent error trends, (3) pedagogical rationale (e.g., reinforcement, difficulty calibration, or cross-domain recovery), and (4) expected cognitive exertion. This transparency builds user trust and fosters metacognitive awareness."
    )

    add_academic_heading(doc, "3.5 Typical End-to-End User & Agent Workflow", level=2)
    add_body_p(
        doc,
        "The standard end-to-end operational cycle follows a structured sequence: (1) User authenticates and selects an operational mode; (2) The agent reads the user's historical cognitive profile and computes target difficulty; (3) The challenge generation engine synthesizes a unique problem with accompanying visual vector parameters; (4) The interactive arena renders the challenge with specialized interaction widgets; (5) The user submits their response while background telemetry captures completion latency; (6) The evaluation module assesses accuracy, classifies mistakes, and updates Elo ratings; (7) Detailed diagnostic feedback is displayed in the review desk; (8) The adaptive session manager updates its pedagogical plan for the subsequent challenge."
    )

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 4: AI & NLP PIPELINE ARCHITECTURE (with Colored Images)
    # =========================================================================
    add_academic_heading(doc, "CHAPTER 4\nAI & NLP PIPELINE ARCHITECTURE", level=1)
    
    add_academic_heading(doc, "4.1 Pipeline Overview", level=2)
    add_body_p(
        doc,
        "The processing of a user interaction in Cognivista follows a seven-stage asynchronous pipeline, illustrated in the architectural diagram in Figure 4.1. Stages 1 through 4 handle intent ingestion, NLP parsing, generative synthesis, and vector schematic construction. Stages 5 through 7 handle attempt capture, diagnostic evaluation, psychometric rating updates, and adaptive curriculum steering."
    )

    # Insert Colored PNG for Figure 4.1
    add_figure_image(doc, "fig4_1_pipeline.png", "Figure 4.1: End-to-End AI & NLP Challenge Processing Pipeline", width_inches=5.4)

    add_academic_heading(doc, "4.2 Pipeline Stage Description", level=2)
    add_body_p(
        doc,
        "Each stage in the seven-tier pipeline executes discrete data transformations with guaranteed validation contracts, as detailed in Table 4.1."
    )

    # Table 4.1
    p_t41 = doc.add_paragraph()
    p_t41.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t41.paragraph_format.space_before = Pt(6)
    p_t41.paragraph_format.space_after = Pt(4)
    r = p_t41.add_run("Table 4.1: Operational Stages of the Cognivista Challenge Pipeline")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t41 = doc.add_table(rows=8, cols=4)
    t41.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t41)
    
    t41_headers = ["Stage", "Input Data", "Algorithmic Processing", "Output Artifact"]
    for i, h in enumerate(t41_headers):
        t41.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t41.rows[0], is_header=True)
    
    t41_data = [
        ("1. Ingestion", "User raw text / UI trigger", "Sanitizes input, detects session state, retrieves user profile", "Sanitized prompt request"),
        ("2. NLP Parsing", "Sanitized prompt request", "Applies regex slot extractors; resolves domain, difficulty, format", "Structured intent payload"),
        ("3. Synthesis", "Structured intent payload", "Orchestrates LLM prompt or curriculum generator for problem generation", "Raw challenge JSON payload"),
        ("4. Vector Gen", "Visual metadata payload", "Computes geometric nodes, circuit wires, gear teeth, balance angles", "Renderable vector SVG schema"),
        ("5. Telemetry", "User arena interaction", "Monitors interaction events, tracks hint usage, computes latency", "User attempt submission"),
        ("6. Diagnosis", "Attempt submission & key", "Executes semantic comparison, classifies error mode, writes feedback", "Diagnostic evaluation report"),
        ("7. Adaptation", "Evaluation report & profile", "Calculates new Elo rating, logs agent decision, adjusts difficulty", "Updated cognitive profile")
    ]
    
    t41_widths = [Inches(1.1), Inches(1.6), Inches(2.2), Inches(1.5)]
    for r_idx, row_data in enumerate(t41_data, start=1):
        row = t41.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t41_widths[i]

    add_academic_heading(doc, "4.3 Layered System Architecture", level=2)
    add_body_p(
        doc,
        "Cognivista follows a clean four-layer decoupled architecture ensuring modularity, horizontal scalability, and testability, as visualized in the full-color system architecture diagram in Figure 4.2."
    )

    # Insert Colored PNG for Figure 4.2
    add_figure_image(doc, "fig4_2_architecture.png", "Figure 4.2: Layered System Architecture of Cognivista", width_inches=5.5)

    add_academic_heading(doc, "4.4 Execution Modes", level=2)
    add_body_p(
        doc,
        "To support diverse learning paradigms, Cognivista provides four distinct operational modes:"
    )
    add_bullet_p(doc, " Self-directed practice across any of the 12 challenge genres with instant feedback and visual schematics.", bold_title="1. Interactive Arena Mode:")
    add_bullet_p(doc, " Freeform exploration allowing users to prompt the agent with natural language to synthesize custom brain challenges on the fly.", bold_title="2. Autonomous Creator Desk:")
    add_bullet_p(doc, " An automated workout where the AI agent steers the curriculum, modulates difficulty, and logs its decisions in real time.", bold_title="3. Adaptive Training Session:")
    add_bullet_p(doc, " Real-world analytical reasoning puzzles built from live external weather, time, and market feeds.", bold_title="4. Live Grounded Intelligence Mode:")

    add_academic_heading(doc, "4.5 Design Principles", level=2)
    add_body_p(
        doc,
        "The system adheres to four foundational engineering principles: (1) Pedagogical Determinism—challenges always maintain verifiable mathematical integrity; (2) Explainable Socratic Feedback—the agent guides learners toward understanding rather than merely displaying answers; (3) Visual-Spatial Primacy—abstract logic is grounded in visual representations wherever feasible; (4) Data Isolation and Privacy—all cognitive telemetry is securely bound to authenticated user accounts."
    )

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 5: MATHEMATICAL FORMULATIONS & ALGORITHMS (SENTENCES, NO ASH BOXES)
    # =========================================================================
    add_academic_heading(doc, "CHAPTER 5\nMATHEMATICAL FORMULATIONS & ALGORITHMS", level=1)
    
    add_body_p(
        doc,
        "Cognivista employs precise mathematical models and psychometric formulations to quantify cognitive competence, modulate challenge complexity, and detect conceptual errors. This chapter articulates the core mathematical formulations and algorithmic procedures underpinning the platform."
    )

    add_academic_heading(doc, "5.1 Elo Cognitive Proficiency Rating Formulation", level=2)
    add_body_p(
        doc,
        "Learner competence is modeled using an adaptation of the Elo rating system, extended to multi-domain psychometrics. Let R_u denote the learner's current rating in domain d, and R_c denote the difficulty rating of the challenge. The expected probability E of the user successfully solving the challenge is defined as:"
    )
    add_body_p(
        doc,
        "E(u, c) = 1 / ( 1 + 10 ^ ( ( R_c - R_u ) / 400 ) )",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold_prefix="Expected Score Formulation:  "
    )
    add_body_p(
        doc,
        "Upon submission, the actual outcome S is recorded as S = 1 for a correct solution and S = 0 for an incorrect solution. The updated rating R'_u is computed as:"
    )
    add_body_p(
        doc,
        "R'_u = R_u + K_u * ( S - E(u, c) )",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold_prefix="Rating Update Formulation:  "
    )
    add_body_p(
        doc,
        "To ensure rapid convergence for novices while preserving rating stability for experienced users, the sensitivity factor K_u is dynamically scaled based on the user's total completed challenges N_u: K_u = 32 if N_u < 30 else (24 if N_u < 100 else 16)."
    )

    add_academic_heading(doc, "5.2 Response Time & Attention Factor", level=2)
    add_body_p(
        doc,
        "Accuracy alone is insufficient to evaluate cognitive fluency. Let T_actual be the time in seconds taken to solve the challenge, and T_expected be the benchmark duration. The Attention & Processing Factor A_f is formulated as:"
    )
    add_body_p(
        doc,
        "A_f = clamp( T_expected / T_actual, 0.5, 2.0 )",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold_prefix="Attention Velocity Ratio:  "
    )

    add_academic_heading(doc, "5.3 Challenge Complexity & Discrimination Index", level=2)
    add_body_p(
        doc,
        "Challenge difficulty R_c is deterministically parameterized by structural complexity metrics: number of reasoning steps N_steps, working memory load M_load, and distractor plausibility D_plaus: R_c = 800 + (250 * N_steps) + (150 * M_load) + (100 * D_plaus)."
    )

    add_academic_heading(doc, "5.4 Adaptive Difficulty Adjustment Function", level=2)
    add_body_p(
        doc,
        "During autonomous training sessions, target difficulty Delta_D for step t+1 is regulated by the agent according to moving average accuracy over the previous w = 3 attempts: Delta_D = +50 if Mean(S_{t-2..t}) > 0.85 else (-50 if Mean(S_{t-2..t}) < 0.40 else 0)."
    )

    add_academic_heading(doc, "5.5 Semantic Intent Similarity & Confidence Metric", level=2)
    add_body_p(
        doc,
        "For natural-language challenge prompts, semantic intent confidence C_intent is computed across keyword overlap and token Jaccard similarity against domain lexicons L_d: Jaccard(P, L_d) = | Tokens(P) intersect L_d | / | Tokens(P) union L_d |."
    )

    add_academic_heading(doc, "5.6 Diagnostic Error Scoring & Misconception Probability", level=2)
    add_body_p(
        doc,
        "When an attempt fails (S = 0), the discrepancy delta = | Answer_user - Answer_correct | is mapped against pre-computed cognitive trap distractor sets D_traps. The misconception probability P(M_k | delta) determines whether the failure stemmed from calculation slippage, misread premise, or spatial inversion."
    )

    add_academic_heading(doc, "5.7 Cognitive Fatigue & Session Progress Calculation", level=2)
    add_body_p(
        doc,
        "Cognitive fatigue F_t is modeled as a leaky accumulator over session steps: F_t = gamma * F_{t-1} + ( T_actual / T_expected ). When F_t exceeds threshold theta_fatigue = 4.5, the agent shifts from complex deductive challenges to restorative attention or memory recall puzzles."
    )

    add_academic_heading(doc, "5.8 Core Algorithms (Procedural Formulations in Sentences)", level=2)
    add_body_p(
        doc,
        "The operational lifecycle of Cognivista is governed by four core algorithms, articulated below in structured academic procedural form:"
    )

    # Algorithm 1 in sentences
    add_body_p(doc, "Algorithm 1 – Autonomous Intent-Driven Challenge Generation", bold_prefix="• ")
    add_algorithm_step(doc, "Step 1:", "The system tokenizes the user's natural language input prompt, strips extraneous punctuation, and converts all tokens to lowercase characters.")
    add_algorithm_step(doc, "Step 2:", "For each supported cognitive domain in the knowledge base, the intent parser computes the token Jaccard similarity and evaluates regular expression pattern match density.")
    add_algorithm_step(doc, "Step 3:", "The target cognitive domain is assigned to the highest-scoring candidate; if scores are equal or zero, the system defaults to the learner's recommended developmental focus.")
    add_algorithm_step(doc, "Step 4:", "When visual keywords (such as angle, triangle, circuit, gear, scale) or spatial domains are detected, the system routes generation to the procedural vector schematic engine; otherwise, a structured few-shot prompt is dispatched to the Gemini 2.5 Flash model.")
    add_algorithm_step(doc, "Step 5:", "The resulting challenge payload is validated against strict Pydantic schemas, confirming that the ground-truth answer and plausible distractor choices are mathematically consistent.")
    add_algorithm_step(doc, "Step 6:", "The fully validated challenge object is returned to the arena execution coordinator for client presentation.")

    # Algorithm 2 in sentences
    add_body_p(doc, "Algorithm 2 – Dynamic Vector Visual Schematic Generation", bold_prefix="• ")
    add_algorithm_step(doc, "Step 1:", "The system reads the visual challenge type and corresponding difficulty rating parameter.")
    add_algorithm_step(doc, "Step 2:", "For geometric triangle problems, it selects random interior angles summing to 180 degrees, assigns dynamic vertex coordinates, and marks one vertex as the unknown target variable x.")
    add_algorithm_step(doc, "Step 3:", "For electric circuit problems, it constructs schematic node coordinates for the battery, series or parallel wire paths, toggle switches, and lightbulb indicator, evaluating electrical continuity.")
    add_algorithm_step(doc, "Step 4:", "For mechanical gear systems, it computes gear radii, tooth counts, and rotational directions such that each subsequent meshed cog alternates between clockwise and counter-clockwise rotation.")
    add_algorithm_step(doc, "Step 5:", "For balance scales, it balances weights across left and right pans, hiding an unknown weight block.")
    add_algorithm_step(doc, "Step 6:", "All calculated geometric coordinates, labels, and vector path definitions are serialized into structured JSON for client-side SVG rendering.")

    # Algorithm 3 in sentences
    add_body_p(doc, "Algorithm 3 – Attempt Evaluation & Diagnostic Mistake Explanation", bold_prefix="• ")
    add_algorithm_step(doc, "Step 1:", "The evaluation engine normalizes the user's submitted response and performs a fuzzy semantic comparison against the ground-truth answer and acceptable synonymous variants.")
    add_algorithm_step(doc, "Step 2:", "If the response matches the correct answer, the system sets the success flag, calculates the positive Elo rating gain using the dynamic sensitivity factor, and prepares affirmative feedback.")
    add_algorithm_step(doc, "Step 3:", "If the response is incorrect, the engine compares the submitted choice against the pre-compiled distractor models to identify the specific cognitive failure mode.")
    add_algorithm_step(doc, "Step 4:", "The failure is categorized according to the cognitive error taxonomy (Calculation Slippage, Misread Premise, Logical Fallacy, Spatial Inversion, or Distractor Trap).")
    add_algorithm_step(doc, "Step 5:", "A targeted Socratic explanation is generated explaining the conceptual flaw without prematurely revealing the solution.")
    add_algorithm_step(doc, "Step 6:", "The complete diagnostic telemetry record is saved to the database and returned to the review desk.")

    # Algorithm 4 in sentences
    add_body_p(doc, "Algorithm 4 – Real-Time Adaptive Training Loop with Agent Decision Stream", bold_prefix="• ")
    add_algorithm_step(doc, "Step 1:", "The autonomous session manager initializes a multi-step training session and retrieves the learner's longitudinal cognitive profile.")
    add_algorithm_step(doc, "Step 2:", "At each step, the agent analyzes recent attempt performance and calculates the moving-average success rate and cognitive fatigue index.")
    add_algorithm_step(doc, "Step 3:", "The agent decides whether to increase difficulty, maintain current difficulty, or switch to a remedial cognitive domain, appending a timestamped entry to the public decision stream log.")
    add_algorithm_step(doc, "Step 4:", "The selected challenge is synthesized and dispatched to the learner's active arena.")
    add_algorithm_step(doc, "Step 5:", "Upon receiving the user's submission, the agent evaluates correctness, measures latency, and adjusts the learner's domain-specific Elo rating.")
    add_algorithm_step(doc, "Step 6:", "Steps 2 through 5 repeat until the total planned workout steps are concluded, at which point a comprehensive session summary report is finalized.")

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 6: TECH STACK & PACKAGES
    # =========================================================================
    add_academic_heading(doc, "CHAPTER 6\nTECH STACK & PACKAGES", level=1)
    
    add_academic_heading(doc, "6.1 Overview", level=2)
    add_body_p(
        doc,
        "Cognivista is engineered as a modern, decoupled web application. The frontend leverages React 19, TypeScript, and Vite for fast client rendering and vector graphics manipulation. The backend is implemented in Python 3.11 with FastAPI for asynchronous REST endpoint handling, Pydantic v2 for data validation, and SQLAlchemy ORM for database operations. Machine learning and NLP capabilities combine the official Google GenAI SDK (Gemini 2.5 Flash) with deterministic regex tokenizers and psychometric scoring libraries."
    )

    add_academic_heading(doc, "6.2 Front-End Technologies", level=2)
    add_bullet_p(doc, " Modern, performant component-based user interface architecture utilizing React hooks for local state and responsive re-rendering.", bold_title="• React 19 & TypeScript 5:")
    add_bullet_p(doc, " High-velocity frontend build pipeline and development server with Hot Module Replacement (HMR).", bold_title="• Vite 6.0:")
    add_bullet_p(doc, " Clean, responsive layout styling using custom CSS variables, dark-mode themes, and subtle micro-animations.", bold_title="• Tailwind CSS 4 & Vanilla Tokens:")
    add_bullet_p(doc, " Vector-based SVG manipulation engine rendering dynamic triangles, electric circuits, rotating gears, and balance scales.", bold_title="• Procedural Vector SVG Engine:")
    add_bullet_p(doc, " Accessible, consistent UI iconography across all dashboards, arenas, and review desks.", bold_title="• Lucide React Icons:")

    add_academic_heading(doc, "6.3 Back-End Technologies", level=2)
    add_bullet_p(doc, " Asynchronous, high-performance Python web framework featuring auto-generated OpenAPI documentation.", bold_title="• FastAPI & Uvicorn:")
    add_bullet_p(doc, " Type-safe data validation, request parsing, and schema generation for challenge JSON objects.", bold_title="• Pydantic v2:")
    add_bullet_p(doc, " Enterprise object-relational mapping supporting SQLite for local zero-config testing and PostgreSQL for production.", bold_title="• SQLAlchemy ORM:")
    add_bullet_p(doc, " Secure cryptographic user credential management utilizing PBKDF2 hashing and JWT bearer tokens.", bold_title="• Passlib & Python-Jose:")

    add_academic_heading(doc, "6.4 AI, NLP, and External Data Libraries", level=2)
    add_bullet_p(doc, " Google GenAI SDK connecting to Gemini 2.5 Flash for rapid few-shot challenge synthesis and Socratic error explanations.", bold_title="• Google GenAI SDK (Gemini 2.5 Flash):")
    add_bullet_p(doc, " Fast, deterministic pattern-matching engine for intent classification, slot filling, and syntax extraction.", bold_title="• Native Python Regex & Tokenizers:")
    add_bullet_p(doc, " Asynchronous HTTP client fetching real-time weather metrics, world clock timezones, and economic indicators.", bold_title="• HTTPX & Requests:")

    # Table 6.1
    p_t61 = doc.add_paragraph()
    p_t61.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t61.paragraph_format.space_before = Pt(6)
    p_t61.paragraph_format.space_after = Pt(4)
    r = p_t61.add_run("Table 6.1: Core Packages and Dependencies Utilized in Cognivista")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t61 = doc.add_table(rows=12, cols=3)
    t61.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t61)
    
    t61_headers = ["Package Name", "Version", "Operational Purpose in Cognivista"]
    for i, h in enumerate(t61_headers):
        t61.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t61.rows[0], is_header=True)
    
    t61_data = [
        ("fastapi", "^0.110.0", "Asynchronous REST API routing, exception handling, and middleware"),
        ("uvicorn", "^0.29.0", "ASGI web server implementation with asynchronous worker pooling"),
        ("pydantic", "^2.6.0", "Strict JSON schema definitions, input validation, and serialization"),
        ("sqlalchemy", "^2.0.28", "Object-relational mapping, database session pooling, and queries"),
        ("google-genai", "^0.1.1", "Official SDK for Gemini 2.5 Flash LLM generative challenge synthesis"),
        ("python-jose", "^3.3.0", "JWT token encoding, decoding, and bearer authentication"),
        ("passlib[bcrypt]", "^1.7.4", "Secure password hashing and cryptographic authentication"),
        ("httpx", "^0.27.0", "Asynchronous HTTP client for live Open-Meteo and WorldTime API feeds"),
        ("react / react-dom", "^19.0.0", "Component hierarchy, virtual DOM reconciliation, and UI lifecycle"),
        ("typescript", "^5.7.0", "Static type checking, interfaces, and compile-time code verification"),
        ("vite", "^6.0.0", "Next-generation frontend tooling, bundling, and local HMR dev server")
    ]
    
    t61_widths = [Inches(1.8), Inches(1.3), Inches(3.1)]
    for r_idx, row_data in enumerate(t61_data, start=1):
        row = t61.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t61_widths[i]

    add_academic_heading(doc, "6.5 Development Environment & Specifications", level=2)
    add_body_p(
        doc,
        "The development and runtime environments for the project are summarized in Table 6.2."
    )

    # Table 6.2
    p_t62 = doc.add_paragraph()
    p_t62.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t62.paragraph_format.space_before = Pt(6)
    p_t62.paragraph_format.space_after = Pt(4)
    r = p_t62.add_run("Table 6.2: System Specifications & Development Environment")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t62 = doc.add_table(rows=7, cols=2)
    t62.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t62)
    
    t62_headers = ["Environment Parameter", "Specification / Configuration Details"]
    for i, h in enumerate(t62_headers):
        t62.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t62.rows[0], is_header=True)
    
    t62_data = [
        ("Host Operating System", "Microsoft Windows 11 Pro / Enterprise (x64)"),
        ("Backend Runtime", "Python 3.11.8 (64-bit virtual environment)"),
        ("Frontend Runtime", "Node.js v20.x with npm package manager"),
        ("Local Port Allocations", "Backend: http://127.0.0.1:8000 | Frontend: http://127.0.0.1:5173"),
        ("Persistent Storage", "SQLite 3 database (cognivista.db) with WAL logging; PostgreSQL ready"),
        ("Supported Browsers", "Google Chrome 120+, Microsoft Edge 120+, Mozilla Firefox 120+")
    ]
    
    t62_widths = [Inches(2.4), Inches(3.8)]
    for r_idx, row_data in enumerate(t62_data, start=1):
        row = t62.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t62_widths[i]

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 7: DATABASE & KNOWLEDGE BASE STRUCTURES (with Colored Diagram)
    # =========================================================================
    add_academic_heading(doc, "CHAPTER 7\nDATABASE & KNOWLEDGE BASE STRUCTURES", level=1)
    
    add_academic_heading(doc, "7.1 Storage Approach & Relational Architecture", level=2)
    add_body_p(
        doc,
        "Cognivista employs an enterprise relational database schema managed through SQLAlchemy ORM. The architecture features full relational referential integrity, foreign key cascades, and indexing on frequently queried columns. For seamless local execution, the system operates against SQLite (cognivista.db), while maintaining complete configuration parity for production PostgreSQL deployments. Every user, challenge item, attempt telemetry record, cognitive domain rating, agent memory, and autonomous decision step is modeled as a dedicated entity, as visualized in the full-color entity-relationship diagram in Figure 7.1."
    )

    # Insert Colored PNG for Figure 7.1
    add_figure_image(doc, "fig7_1_class_diagram.png", "Figure 7.1: Entity-Relationship & Relational Data Model of Cognivista", width_inches=5.6)

    add_academic_heading(doc, "7.2 User Entity (users Table)", level=2)
    add_body_p(
        doc,
        "The User entity stores authenticated account credentials, security flags, and client visual preferences, as defined in Table 7.1."
    )

    # Table 7.1
    p_t71 = doc.add_paragraph()
    p_t71.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t71.paragraph_format.space_before = Pt(6)
    p_t71.paragraph_format.space_after = Pt(4)
    r = p_t71.add_run("Table 7.1: User Entity Schema (users Table)")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t71 = doc.add_table(rows=8, cols=4)
    t71.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t71)
    
    t71_headers = ["Field Name", "Data Type", "Constraints", "Description"]
    for i, h in enumerate(t71_headers):
        t71.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t71.rows[0], is_header=True)
    
    t71_data = [
        ("id", "String(36)", "PRIMARY KEY, UUID", "Unique identifier for the user account"),
        ("email", "String(255)", "UNIQUE, INDEX, NOT NULL", "User email address used for login and notifications"),
        ("username", "String(100)", "UNIQUE, INDEX, NOT NULL", "Display username and unique handle across the platform"),
        ("full_name", "String(255)", "NULLABLE", "User's real or academic display name"),
        ("hashed_password", "String(255)", "NOT NULL", "PBKDF2/Bcrypt cryptographically salted password hash"),
        ("is_active", "Boolean", "DEFAULT True", "Account status flag indicating active or suspended state"),
        ("settings", "JSON", "DEFAULT SettingsDict", "User UI preferences (sound, haptics, theme, difficulty mode)")
    ]
    
    t71_widths = [Inches(1.5), Inches(1.3), Inches(1.6), Inches(2.0)]
    for r_idx, row_data in enumerate(t71_data, start=1):
        row = t71.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t71_widths[i]

    add_academic_heading(doc, "7.3 Challenge Entity (challenges Table)", level=2)
    add_body_p(
        doc,
        "The Challenge entity stores dynamically generated and curated cognitive puzzles, including visual SVG payload data, as detailed in Table 7.2."
    )

    # Table 7.2
    p_t72 = doc.add_paragraph()
    p_t72.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t72.paragraph_format.space_before = Pt(6)
    p_t72.paragraph_format.space_after = Pt(4)
    r = p_t72.add_run("Table 7.2: Challenge Entity Schema (challenges Table)")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t72 = doc.add_table(rows=11, cols=4)
    t72.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t72)
    
    for i, h in enumerate(t71_headers):
        t72.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t72.rows[0], is_header=True)
    
    t72_data = [
        ("id", "String(36)", "PRIMARY KEY, UUID", "Unique identifier for the synthesized challenge"),
        ("challenge_type", "String(50)", "INDEX, NOT NULL", "Specific category: spatial_rotation, logic_deduction, circuit, etc."),
        ("primary_domain", "String(50)", "INDEX, NOT NULL", "Core faculty: logic, numerical, spatial, memory, verbal, etc."),
        ("title", "String(255)", "NOT NULL", "Brief, engaging title describing the challenge"),
        ("prompt", "Text", "NOT NULL", "Full problem statement presented to the learner"),
        ("payload", "JSON", "NOT NULL", "Interactive parameters: choices, SVG coordinates, gear counts"),
        ("correct_answer", "JSON", "NOT NULL", "Ground-truth answer value or key-value map"),
        ("explanation", "Text", "NOT NULL", "Detailed pedagogical solution breakdown"),
        ("difficulty_rating", "Float", "DEFAULT 1200.0", "Item psychometric difficulty rating on the Elo scale"),
        ("source_info", "JSON", "DEFAULT LiveDict", "Attribution data for real-world live challenges (API, timestamp)")
    ]
    
    for r_idx, row_data in enumerate(t72_data, start=1):
        row = t72.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t71_widths[i]

    add_academic_heading(doc, "7.4 Attempt Entity (attempts Table)", level=2)
    add_body_p(
        doc,
        "The Attempt entity logs every user interaction, measuring response time, accuracy, and diagnostic mistake classifications, as specified in Table 7.3."
    )

    # Table 7.3
    p_t73 = doc.add_paragraph()
    p_t73.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t73.paragraph_format.space_before = Pt(6)
    p_t73.paragraph_format.space_after = Pt(4)
    r = p_t73.add_run("Table 7.3: Attempt Entity Schema (attempts Table)")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t73 = doc.add_table(rows=10, cols=4)
    t73.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t73)
    
    for i, h in enumerate(t71_headers):
        t73.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t73.rows[0], is_header=True)
    
    t73_data = [
        ("id", "String(36)", "PRIMARY KEY, UUID", "Unique identifier for the attempt record"),
        ("user_id", "String(36)", "FOREIGN KEY, CASCADE", "Reference to the submitting user in users.id"),
        ("challenge_id", "String(36)", "FOREIGN KEY, CASCADE", "Reference to the attempted challenge in challenges.id"),
        ("user_response", "JSON", "NOT NULL", "The exact value, choice ID, or matching pairs submitted"),
        ("is_correct", "Boolean", "NOT NULL", "Evaluation verdict: True if correct, False otherwise"),
        ("response_time_ms", "Integer", "NOT NULL", "Total duration in milliseconds spent on the challenge"),
        ("pre_attempt_elo", "Float", "NOT NULL", "Learner's domain rating before the attempt was evaluated"),
        ("post_attempt_elo", "Float", "NOT NULL", "Learner's updated domain rating following the attempt"),
        ("mistake_analysis", "JSON", "DEFAULT Dict", "Diagnostic taxonomy breakdown: mechanism, deviation, remediation")
    ]
    
    for r_idx, row_data in enumerate(t73_data, start=1):
        row = t73.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t71_widths[i]

    add_academic_heading(doc, "7.5 CognitiveProfile Entity (cognitive_profiles Table)", level=2)
    add_body_p(
        doc,
        "The CognitiveProfile entity maintains persistent psychometric telemetry across seven core cognitive faculties, as outlined in Table 7.4."
    )

    # Table 7.4
    p_t74 = doc.add_paragraph()
    p_t74.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t74.paragraph_format.space_before = Pt(6)
    p_t74.paragraph_format.space_after = Pt(4)
    r = p_t74.add_run("Table 7.4: CognitiveProfile Entity Schema (cognitive_profiles Table)")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t74 = doc.add_table(rows=8, cols=4)
    t74.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t74)
    
    for i, h in enumerate(t71_headers):
        t74.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t74.rows[0], is_header=True)
    
    t74_data = [
        ("id", "String(36)", "PRIMARY KEY, UUID", "Unique profile record identifier"),
        ("user_id", "String(36)", "FOREIGN KEY, UNIQUE", "One-to-one link to owning user in users.id"),
        ("composite_elo", "Float", "DEFAULT 1200.0", "Aggregated general intelligence and cognitive rating"),
        ("domain_ratings", "JSON", "DEFAULT DomainMap", "Fine-grained Elo scores: logic, numerical, spatial, memory, etc."),
        ("domain_stats", "JSON", "DEFAULT StatsMap", "Per-domain counts: solved, attempts, avg_time_ms, hints"),
        ("cognitive_bottlenecks", "JSON", "DEFAULT List", "List of identified conceptual weaknesses diagnosed by the agent"),
        ("accuracy_history", "JSON", "DEFAULT List", "Chronological telemetry log of recent attempt outcomes")
    ]
    
    for r_idx, row_data in enumerate(t74_data, start=1):
        row = t74.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t71_widths[i]

    add_academic_heading(doc, "7.6 AgentMemory & Conversation Entities", level=2)
    add_body_p(
        doc,
        "To support long-term personalization, the AgentMemory entity stores episodic facts, while ConversationMessage stores multi-turn dialogue from the review desk, as described in Table 7.5 and Table 7.6."
    )

    # Table 7.5
    p_t75 = doc.add_paragraph()
    p_t75.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t75.paragraph_format.space_before = Pt(6)
    p_t75.paragraph_format.space_after = Pt(4)
    r = p_t75.add_run("Table 7.5: AgentMemory Entity Schema (agent_memories Table)")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t75 = doc.add_table(rows=6, cols=4)
    t75.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t75)
    
    for i, h in enumerate(t71_headers):
        t75.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t75.rows[0], is_header=True)
    
    t75_data = [
        ("id", "String(36)", "PRIMARY KEY, UUID", "Unique identifier for memory entry"),
        ("user_id", "String(36)", "FOREIGN KEY, CASCADE", "Target user reference in users.id"),
        ("memory_type", "String(50)", "NOT NULL", "Category: cognitive_bottleneck, preference, episodic_fact"),
        ("key", "String(100)", "NOT NULL", "Indexed lookup key for fast memory retrieval"),
        ("value", "JSON", "NOT NULL", "Structured observation payload and agent confidence score")
    ]
    
    for r_idx, row_data in enumerate(t75_data, start=1):
        row = t75.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t71_widths[i]

    # Table 7.6
    p_t76 = doc.add_paragraph()
    p_t76.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t76.paragraph_format.space_before = Pt(6)
    p_t76.paragraph_format.space_after = Pt(4)
    r = p_t76.add_run("Table 7.6: ConversationMessage Entity Schema (conversation_messages Table)")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t76 = doc.add_table(rows=6, cols=4)
    t76.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t76)
    
    for i, h in enumerate(t71_headers):
        t76.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t76.rows[0], is_header=True)
    
    t76_data = [
        ("id", "String(36)", "PRIMARY KEY, UUID", "Unique message identifier"),
        ("user_id", "String(36)", "FOREIGN KEY, CASCADE", "User reference in users.id"),
        ("role", "String(20)", "NOT NULL", "Message sender: 'user', 'assistant', or 'system'"),
        ("content", "Text", "NOT NULL", "Linguistic content of the dialogue turn"),
        ("metadata", "JSON", "DEFAULT Dict", "Linked challenge ID, attempt ID, and pedagogical intents")
    ]
    
    for r_idx, row_data in enumerate(t76_data, start=1):
        row = t76.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t71_widths[i]

    add_academic_heading(doc, "7.7 Knowledge Base Structures & Curated Taxonomies", level=2)
    add_body_p(
        doc,
        "In addition to relational tables, the Cognivista agent relies on structured in-memory knowledge bases. These curated ontologies serve as the cognitive domain model, guiding procedural challenge synthesis and error diagnosis."
    )

    # Table 7.7
    p_t77 = doc.add_paragraph()
    p_t77.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t77.paragraph_format.space_before = Pt(6)
    p_t77.paragraph_format.space_after = Pt(4)
    r = p_t77.add_run("Table 7.7: Dynamic Challenge Universe Taxonomy (12 Cognitive Genres)")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t77 = doc.add_table(rows=13, cols=3)
    t77.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t77)
    
    t77_headers = ["Genre Identifier", "Target Cognitive Faculty", "Interactive Task Paradigm"]
    for i, h in enumerate(t77_headers):
        t77.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t77.rows[0], is_header=True)
    
    t77_data = [
        ("geometry_visual", "Spatial-Geometric Reasoning", "Calculate unknown interior/exterior angles using vector SVG schematics"),
        ("circuit_logic", "Systems Analysis & Boolean Flow", "Determine lightbulb states across parallel/series switch networks"),
        ("gear_rotation", "Mechanical & Kinematic Reasoning", "Infer rotational direction and speed across interlocking gear trains"),
        ("balance_scale", "Algebraic & Quantitative Intuition", "Deduce unknown block weights on two-pan balance scale diagrams"),
        ("fraction_pie", "Proportional & Fractional Modeling", "Evaluate shaded vs. unshaded sectors in radial division graphs"),
        ("flight_trajectory", "Spatiotemporal & Time-Zone Logic", "Calculate flight durations and arrival times across world waypoints"),
        ("logic_deduction", "Formal Deductive Logic", "Evaluate syllogisms, truth-teller puzzles, and grid matrices"),
        ("matching_pairs", "Associative Semantic Memory", "Drag-and-drop connection between complementary cognitive concepts"),
        ("ordering_sequence", "Temporal & Algorithmic Sequencing", "Arrange scrambled historical, procedural, or numeric stages"),
        ("anagram_word", "Lexical Access & Mental Synthesis", "Unscramble letter constellations under time constraints"),
        ("odd_one_out", "Categorical Classification", "Identify semantic or visual outliers within a feature group"),
        ("detective_clues", "Lateral & Inductive Investigation", "Synthesize multi-source forensic clues to identify culprits")
    ]
    
    t77_widths = [Inches(1.8), Inches(2.2), Inches(2.4)]
    for r_idx, row_data in enumerate(t77_data, start=1):
        row = t77.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t77_widths[i]

    # Table 7.8
    p_t78 = doc.add_paragraph()
    p_t78.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t78.paragraph_format.space_before = Pt(8)
    p_t78.paragraph_format.space_after = Pt(4)
    r = p_t78.add_run("Table 7.8: Cognitive Error & Misconception Taxonomy")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t78 = doc.add_table(rows=6, cols=3)
    t78.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t78)
    
    t78_headers = ["Error Classification", "Typical Cognitive Manifestation", "Agent Remediation Strategy"]
    for i, h in enumerate(t78_headers):
        t78.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t78.rows[0], is_header=True)
    
    t78_data = [
        ("Calculation Slippage", "Correct conceptual formula applied but arithmetic mistake in execution", "Highlights exact arithmetic step; prompts mental verification check"),
        ("Misread Premise", "Overlooking critical negative constraint (e.g., 'which is NOT true')", "Displays highlighted excerpt of the overlooked rule or boundary"),
        ("Logical Fallacy", "Affirming the consequent or confusing correlation with causation", "Provides counterexample demonstrating invalidity of the inference"),
        ("Spatial Inversion", "Confusing clockwise/counter-clockwise rotation or mirror axes", "Animates step-by-step rotational trajectory on the vector canvas"),
        ("Distractor Trap", "Selecting an enticing option designed to capture common heuristic biases", "Explains why the trap seemed plausible and contrasts it with the truth")
    ]
    
    for r_idx, row_data in enumerate(t78_data, start=1):
        row = t78.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t77_widths[i]

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 8: MODULES & IMPLEMENTATION WALKTHROUGH (with Colored State Diagram)
    # =========================================================================
    add_academic_heading(doc, "CHAPTER 8\nMODULES & IMPLEMENTATION WALKTHROUGH", level=1)
    
    add_academic_heading(doc, "8.1 Project Directory Structure", level=2)
    add_body_p(
        doc,
        "The complete Cognivista codebase is organized into modular packages across backend and frontend environments:"
    )

    proj_tree = (
        "cognivista/\n"
        "  |-- backend/\n"
        "  |     |-- app/\n"
        "  |     |     |-- main.py                    # FastAPI application entrypoint\n"
        "  |     |     |-- core/database.py           # SQLAlchemy database session & engine\n"
        "  |     |     |-- models/                    # ORM entities (User, Challenge, Attempt...)\n"
        "  |     |     |-- api/                       # REST endpoints (auth, challenges, agent...)\n"
        "  |     |     |-- agent/orchestrator.py      # Autonomous agent workflow orchestrator\n"
        "  |     |     |-- nlp/intent_parser.py       # Regex & neural intent slot extractor\n"
        "  |     |     |-- generators/                # Wayground & genre challenge engines\n"
        "  |     |     |-- evaluator/evaluator.py     # Response verification & error classifier\n"
        "  |     |     +-- services/live_data.py      # Real-world API connectors (Weather, Time)\n"
        "  |     +-- cognivista.db                    # SQLite persistent database\n"
        "  +-- frontend/\n"
        "        |-- src/\n"
        "        |     |-- App.tsx                    # Main container & mode router\n"
        "        |     |-- types/index.ts             # TypeScript domain interfaces\n"
        "        |     +-- components/\n"
        "        |           |-- Arena/VisualChallengeBoard.tsx # Dynamic vector SVG engine\n"
        "        |           |-- Create/CreateChallengeView.tsx # Natural-language creator desk\n"
        "        |           |-- Review/ChallengeReviewView.tsx # Diagnostic review & dialogue\n"
        "        |           |-- Session/TrainingSessionView.tsx# Autonomous session & decision stream\n"
        "        |           +-- Live/LiveChallengesView.tsx    # Grounded environmental feeds\n"
        "        +-- package.json                     # Frontend build configurations"
    )
    p_tree = doc.add_paragraph()
    p_tree.paragraph_format.left_indent = Inches(0.2)
    r_t = p_tree.add_run(proj_tree)
    r_t.font.name = 'Consolas'
    r_t.font.size = Pt(8.5)

    add_academic_heading(doc, "8.2 Module Summary", level=2)
    add_body_p(
        doc,
        "Table 8.1 outlines the primary source code files and their core architectural responsibilities."
    )

    # Table 8.1
    p_t81 = doc.add_paragraph()
    p_t81.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t81.paragraph_format.space_before = Pt(6)
    p_t81.paragraph_format.space_after = Pt(4)
    r = p_t81.add_run("Table 8.1: Core Implementation Modules and Responsibilities")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t81 = doc.add_table(rows=10, cols=3)
    t81.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t81)
    
    t81_headers = ["Module / Source File", "Architectural Layer", "Primary Operational Responsibility"]
    for i, h in enumerate(t81_headers):
        t81.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t81.rows[0], is_header=True)
    
    t81_data = [
        ("app/api/auth.py", "API / Security", "User registration, credential verification, and JWT bearer token issuance"),
        ("app/nlp/intent_parser.py", "NLP / Ingestion", "Regex slot extraction and prompt intent classification across 12 domains"),
        ("app/agent/orchestrator.py", "Agent Core", "Autonomous decision stream logging, skill routing, and training sessions"),
        ("app/generators/wayground_fetcher.py", "Challenge Engine", "Generates authentic visual challenges (triangles, circuits, gears, scales)"),
        ("app/evaluator/evaluator.py", "Evaluation / NLP", "Evaluates responses, computes Elo gains, and classifies error taxonomy"),
        ("app/services/live_data.py", "External Feeds", "Ingests live Open-Meteo, World Time, and market data with attribution"),
        ("src/components/Arena/VisualChallengeBoard.tsx", "Frontend / SVG", "Procedural vector graphics rendering for geometric, circuit, and gear puzzles"),
        ("src/components/Create/CreateChallengeView.tsx", "Frontend / UI", "4-step natural language challenge synthesis stepper and preview desk"),
        ("src/components/Review/ChallengeReviewView.tsx", "Frontend / UI", "Post-attempt review desk and conversational Socratic remediation desk")
    ]
    
    t81_widths = [Inches(2.1), Inches(1.5), Inches(2.8)]
    for r_idx, row_data in enumerate(t81_data, start=1):
        row = t81.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t81_widths[i]

    add_academic_heading(doc, "8.3 Autonomous Training Session State Machine", level=2)
    add_body_p(
        doc,
        "Figure 8.1 illustrates the state transitions executed by the Autonomous Training Session Manager during an active workout."
    )

    # Insert Colored PNG for Figure 8.1
    add_figure_image(doc, "fig8_1_state_machine.png", "Figure 8.1: Autonomous Training Session State Transitions", width_inches=5.4)

    add_academic_heading(doc, "8.4 REST API Endpoints Specification", level=2)
    add_body_p(
        doc,
        "Cognivista exposes a comprehensive suite of REST endpoints, summarized in Table 8.2."
    )

    # Table 8.2
    p_t82 = doc.add_paragraph()
    p_t82.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t82.paragraph_format.space_before = Pt(6)
    p_t82.paragraph_format.space_after = Pt(4)
    r = p_t82.add_run("Table 8.2: Cognivista REST API Endpoints")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    t82 = doc.add_table(rows=9, cols=3)
    t82.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t82)
    
    t82_headers = ["HTTP Method & Route", "Input Payload / Params", "Operational Purpose"]
    for i, h in enumerate(t82_headers):
        t82.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    format_row(t82.rows[0], is_header=True)
    
    t82_data = [
        ("POST /api/auth/login", "OAuth2 Form (username, password)", "Authenticates user and returns JWT bearer token"),
        ("GET /api/profile/history", "Authorization: Bearer <token>", "Retrieves full historical attempts with prompts and diagnoses"),
        ("POST /api/challenges/generate", "JSON: { prompt, domain, difficulty }", "Synthesizes custom challenge via NLP orchestrator"),
        ("GET /api/challenges/wayground", "Query: ?genre=<genre>&difficulty=<elo>", "Fetches authentic Wayground-style challenge with SVG data"),
        ("POST /api/challenges/attempt", "JSON: { challenge_id, user_response, time_ms }", "Evaluates answer, updates Elo ratings, and records attempt"),
        ("POST /api/challenges/attempt/{id}/explain", "JSON: { user_question }", "Conversational Socratic explanation of why answer was wrong"),
        ("POST /api/agent/session/start", "JSON: { target_steps, initial_domains }", "Initializes autonomous adaptive training workout session"),
        ("GET /api/live/intel", "None", "Returns real-time weather, world time, and market intelligence")
    ]
    
    t82_widths = [Inches(2.4), Inches(2.2), Inches(1.8)]
    for r_idx, row_data in enumerate(t82_data, start=1):
        row = t82.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].paragraphs[0].add_run(val)
        format_row(row)
        for i, c in enumerate(row.cells):
            c.width = t82_widths[i]

    add_academic_heading(doc, "8.5 Automated Verification & Testing Suite", level=2)
    add_body_p(
        doc,
        "The application was validated through an automated test suite verifying: (1) API token security, (2) intent parsing across edge-case syntax, (3) SVG coordinate integrity, and (4) Elo rating monotonicity. The frontend build completed cleanly via `tsc -b && vite build` in 3.69s with 0 errors across 1,842 modules."
    )

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 9: OUTPUT SCREENSHOTS (with Real Coloured Images)
    # =========================================================================
    add_academic_heading(doc, "CHAPTER 9\nOUTPUT SCREENSHOTS", level=1)
    
    add_body_p(
        doc,
        "This chapter presents full-color screenshots of the running Cognivista application across its six primary operational modules. All screenshots were captured directly from the executing local deployment."
    )

    add_academic_heading(doc, "9.1 Cognitive Analytics & Progress Tracker Dashboard", level=2)
    add_body_p(
        doc,
        "Figure 9.1 illustrates the Cognitive Progress Dashboard. The interface displays overall user competence, composite Elo ratings, domain-specific radar charts, response latency percentiles, and an interactive workout calendar tracking daily cognitive exertion."
    )
    add_figure_image(doc, "fig9_1_dashboard.png", "Figure 9.1: Cognitive Analytics & Progress Tracker Dashboard", width_inches=5.5)

    add_academic_heading(doc, "9.2 Natural-Language Challenge Creator Desk", level=2)
    add_body_p(
        doc,
        "Figure 9.2 demonstrates the Natural-Language Challenge Creator Desk. The user enters freeform instructions. The 4-step agent reasoning stepper (Intent Analysis -> Domain Routing -> Parameter Synthesis -> Vector Payload Construction) executes live, displaying a real-time challenge preview."
    )
    add_figure_image(doc, "fig9_2_creator_desk.png", "Figure 9.2: Natural-Language Challenge Creator Desk with 4-Step Stepper", width_inches=5.5)

    add_academic_heading(doc, "9.3 Dynamic Challenge Arena — Geometric Angle Solving", level=2)
    add_body_p(
        doc,
        "Figure 9.3 displays the Interactive Challenge Arena rendering a Wayground-style geometric triangle puzzle. The dynamic vector SVG engine draws triangle vertices with labeled interior angles (65°, 50°) and an unknown target (x°), accompanied by an interactive numerical selection keypad."
    )
    add_figure_image(doc, "fig9_3_geometry_arena.png", "Figure 9.3: Dynamic Challenge Arena — Geometric Angle Solving", width_inches=5.5)

    add_academic_heading(doc, "9.4 Dynamic Challenge Arena — Electric Circuit Schematics", level=2)
    add_body_p(
        doc,
        "Figure 9.4 illustrates the Electric Circuit Arena. The vector schematic renders a DC battery source, parallel wiring branches, interactive toggle switches, and a filament indicator, prompting the learner to deduce current flow under varying switch configurations."
    )
    add_figure_image(doc, "fig9_4_circuit_arena.png", "Figure 9.4: Electric Circuit Schematic Challenge Arena", width_inches=5.5)

    add_academic_heading(doc, "9.5 Dynamic Challenge Arena — Interlocking Gear Trains", level=2)
    add_body_p(
        doc,
        "Figure 9.5 showcases the Mechanical Gear Arena. Three intermeshed cogs are rendered with tooth ratios and rotational direction vectors. The learner determines the rotational direction and relative speed of the terminal gear."
    )
    add_figure_image(doc, "fig9_5_gear_arena.png", "Figure 9.5: Interlocking Gear Train & Mechanical Reasoning Arena", width_inches=5.5)

    add_academic_heading(doc, "9.6 Interactive Multi-Arena Interaction Modes", level=2)
    add_body_p(
        doc,
        "Figure 9.6 displays Cognivista's diverse interaction paradigms: drag-and-drop matching, chronological ordering cards, True/False verification switches, anagram unscrambling chips, and detective clue matrices."
    )
    add_figure_image(doc, "fig9_6_multi_arena.png", "Figure 9.6: Interactive Multi-Arena Modes (Matching, Ordering, Anagrams)", width_inches=5.5)

    add_academic_heading(doc, "9.7 Post-Attempt Review Desk & Socratic Remediation Dialogue", level=2)
    add_body_p(
        doc,
        "Figure 9.7 presents the Post-Attempt Review Desk. The interface displays side-by-side comparisons of submitted vs. correct answers, response time metrics, and error classifications. Below, an interactive conversational desk enables learners to ask 'Why was my answer wrong?' to receive Socratic guidance."
    )
    add_figure_image(doc, "fig9_7_review_desk.png", "Figure 9.7: Post-Attempt Review Desk & 'Why was my answer wrong?' Chat", width_inches=5.5)

    add_academic_heading(doc, "9.8 Adaptive Training Session & Live Agent Decision Stream", level=2)
    add_body_p(
        doc,
        "Figure 9.8 shows the Autonomous Adaptive Training Session in progress. The right sidebar features the live AI AGENT AUTONOMOUS DECISION STREAM log, detailing the agent's real-time pedagogical reasoning during challenge selection."
    )
    add_figure_image(doc, "fig9_8_training_session.png", "Figure 9.8: Adaptive Training Session & Live Agent Decision Stream", width_inches=5.5)

    add_academic_heading(doc, "9.9 Real-World Live Data Intelligence Challenges", level=2)
    add_body_p(
        doc,
        "Figure 9.9 illustrates the Live Challenges module. Dynamic analytical puzzles are generated from live Open-Meteo weather readings, world timezone clocks, and economic indices, displaying verified source attribution and freshness timestamps."
    )
    add_figure_image(doc, "fig9_9_live_challenges.png", "Figure 9.9: Real-World Live Data Intelligence Challenges Desk", width_inches=5.5)

    doc.add_page_break()

    # =========================================================================
    # CHAPTER 10: CONCLUSION & FUTURE ENHANCEMENTS
    # =========================================================================
    add_academic_heading(doc, "CHAPTER 10\nCONCLUSION & FUTURE ENHANCEMENTS", level=1)
    
    add_academic_heading(doc, "10.1 Conclusion", level=2)
    add_body_p(
        doc,
        "The Cognivista project successfully demonstrates that uniting modern Natural Language Processing, generative LLM orchestration, procedural vector rendering, and psychometric Elo modeling produces an adaptive, engaging, and pedagogically robust brain-training platform. By liberating cognitive challenges from static, hardcoded databases, Cognivista provides an infinite universe of diverse puzzles spanning 12 distinct intellectual domains."
    )
    add_body_p(
        doc,
        "The introduction of procedural vector SVG schematics grounds abstract logical and spatial problems in authentic visual manipulation, addressing the limitations of text-only brain-training software. Furthermore, the hybrid response evaluation engine provides explainable error classification and interactive Socratic remediation, empowering learners to understand the cognitive mechanisms underlying their mistakes. The transparent Autonomous Decision Stream bridges the gap between learner and AI, delivering a personalized training regime that continuously adapts to individual growth."
    )

    add_academic_heading(doc, "10.2 System Strengths Observed", level=2)
    add_bullet_p(doc, " The system generates unique, pedagogically sound puzzles on demand without relying on repetitive question pools.", bold_title="• Endless Procedural Challenge Generation:")
    add_bullet_p(doc, " Mathematical precision ensures that geometric angles, electric switches, and gear ratios are always solvable and visually clear.", bold_title="• Mathematical & Visual Integrity:")
    add_bullet_p(doc, " Rather than opaque scoring, every error is classified into pedagogical categories, supported by interactive conversational remediation.", bold_title="• Diagnostic Socratic Explainability:")
    add_bullet_p(doc, " Dynamic K-factor Elo ratings accurately quantify competence across distinct faculties, preventing rating distortion from lucky guesses.", bold_title="• Multi-Domain Psychometric Modeling:")
    add_bullet_p(doc, " Real-world environmental grounding bridges the gap between theoretical puzzles and practical analytical reasoning.", bold_title="• Live Physical World Grounding:")

    add_academic_heading(doc, "10.3 Limitations Observed", level=2)
    add_bullet_p(doc, " While regex and few-shot templates capture intent reliably, complex syntactic compound sentences may occasionally fall back to primary domain defaults.", bold_title="• Compound Syntax Ambiguity:")
    add_bullet_p(doc, " Generating novel multi-modal SVG schematics entirely through LLMs can occasionally cause minor visual overlap, which is why Cognivista utilizes deterministic parametric templates for core visual genres.", bold_title="• Procedural Vector Constraints:")
    add_bullet_p(doc, " High-difficulty lateral thinking scenarios with ambiguous interpretations require extensive semantic synonym mapping to eliminate false-negative evaluations.", bold_title="• Lateral Scenario Disambiguation:")

    add_academic_heading(doc, "10.4 Future Enhancements", level=2)
    add_bullet_p(doc, " Integrate low-latency eye-tracking via standard webcams and consumer EEG telemetry to directly measure cognitive load and attentional focus during problem solving.", bold_title="• Multi-Modal Eye-Tracking & EEG Telemetry:")
    add_bullet_p(doc, " Expand the vector rendering pipeline to support WebGL-powered 3D mental rotation and physics-based fluid simulation puzzles.", bold_title="• 3D Spatial & Physics Simulation Engine:")
    add_bullet_p(doc, " Enable real-time asynchronous multiplayer brain challenge duels where participants solve procedurally balanced puzzles concurrently.", bold_title="• Collaborative & Competitive Multi-Agent Duels:")
    add_bullet_p(doc, " Extend native mobile applications utilizing React Native with offline challenge caching and push-notification workout reminders.", bold_title="• Cross-Platform Mobile Deployment:")
    add_bullet_p(doc, " Integrate with academic institution Learning Management Systems (Canvas, Moodle) via LTI standards for automated grading and cognitive analytics reporting.", bold_title="• Academic LMS Integration:")

    doc.add_page_break()

    # =========================================================================
    # REFERENCES
    # =========================================================================
    add_academic_heading(doc, "REFERENCES", level=1)
    
    refs = [
        "[1] J. R. Anderson, Rules of the Mind. Hillsdale, NJ: Lawrence Erlbaum Associates, 1993.",
        "[2] B. F. Skinner, The Technology of Teaching. New York: Appleton-Century-Crofts, 1968.",
        "[3] A. Elo, The Rating of Chessplayers, Past and Present. New York: Arco Publishing, 1978.",
        "[4] A. Vaswani et al., \"Attention is all you need,\" in Advances in Neural Information Processing Systems (NeurIPS), 2017, pp. 5998–6008.",
        "[5] T. B. Brown et al., \"Language models are few-shot learners,\" in Advances in Neural Information Processing Systems (NeurIPS), 2020, pp. 1877–1901.",
        "[6] J. Sweller, \"Cognitive load during problem solving: Effects on learning,\" Cognitive Science, vol. 12, no. 2, pp. 257–285, 1988.",
        "[7] K. VanLehn, \"The relative effectiveness of human tutoring, intelligent tutoring systems, and other tutoring systems,\" Educational Psychologist, vol. 46, no. 4, pp. 197–221, 2011.",
        "[8] S. E. Embretson and S. P. Reise, Item Response Theory for Psychologists. Mahwah, NJ: Lawrence Erlbaum Associates, 2000.",
        "[9] D. Corbett, Cognitive Systems and Artificial Intelligence. Berlin: Springer, 2018.",
        "[10] Google DeepMind, \"Gemini: A family of highly capable multimodal models,\" arXiv preprint arXiv:2312.11805, 2023.",
        "[11] M. Chi, \"Active-constructive-interactive: A conceptual framework for differentiating learning activities,\" Topics in Cognitive Science, vol. 1, no. 1, pp. 73–105, 2009.",
        "[12] J. Piaget, The Psychology of Intelligence. London: Routledge, 1950."
    ]
    for r_entry in refs:
        add_body_p(doc, r_entry, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)

    # Save to COGNIVISTA_PROJECT_REPORT_V4.docx
    v4_path = os.path.abspath("COGNIVISTA_PROJECT_REPORT_V4.docx")
    doc.save(v4_path)
    print(f"Full perfect academic report successfully saved at: {v4_path}")

    # Also attempt saving to COGNIVISTA_PROJECT_REPORT_UPDATED.docx and COGNIVISTA_PROJECT_REPORT.docx if unlocked
    for alt in ["COGNIVISTA_PROJECT_REPORT_UPDATED.docx", "COGNIVISTA_PROJECT_REPORT_FINAL.docx", "COGNIVISTA_PROJECT_REPORT.docx"]:
        try:
            doc.save(os.path.abspath(alt))
            print(f"Also saved successfully to: {alt}")
        except PermissionError:
            print(f"{alt} is currently open in Word; user can open {v4_path}.")

if __name__ == "__main__":
    build_perfect_report()
