import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Sets internal padding for a cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Sets neat academic borders for a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_box_borders(table, color="000000", sz="6"):
    """Sets a complete outer border for diagrams/code blocks/placeholders."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_row(row, is_header=False):
    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        if is_header:
            set_cell_background(cell, "EAEAEA")

def add_academic_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.color.rgb = RGBColor(0, 0, 0)
    
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(12)
        r.font.size = Pt(14)
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(12.5)
    elif level == 3:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(12)
    return p

def add_body_p(doc, text, bold_prefix=None, italic_prefix=None, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.bold = True
        rb.font.name = 'Times New Roman'
        rb.font.size = Pt(12)
        rb.font.color.rgb = RGBColor(0, 0, 0)
        
    if italic_prefix:
        ri = p.add_run(italic_prefix)
        ri.italic = True
        ri.font.name = 'Times New Roman'
        ri.font.size = Pt(12)
        ri.font.color.rgb = RGBColor(0, 0, 0)
        
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet_p(doc, text, bold_title=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    if bold_title:
        rb = p.add_run(bold_title)
        rb.bold = True
        rb.font.name = 'Times New Roman'
        rb.font.size = Pt(12)
        rb.font.color.rgb = RGBColor(0, 0, 0)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_code_box(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_box_borders(tbl, color="CCCCCC", sz="4")
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F8F8")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code_text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

def add_diagram_box(doc, diagram_text, figure_caption):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_box_borders(tbl, color="888888", sz="6")
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "FAFAFA")
    set_cell_margins(cell, top=160, bottom=160, left=180, right=180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(diagram_text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(6)
    p_cap.paragraph_format.space_after = Pt(12)
    rc = p_cap.add_run(figure_caption)
    rc.bold = True
    rc.font.name = 'Times New Roman'
    rc.font.size = Pt(11)
    rc.font.color.rgb = RGBColor(0, 0, 0)

def add_screenshot_placeholder(doc, placeholder_title, figure_caption, height_in_inches=3.0):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_box_borders(tbl, color="999999", sz="6")
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F5F5F5")
    set_cell_margins(cell, top=180, bottom=180, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f"[ {placeholder_title} ]\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    r.bold = True
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    r2 = p.add_run("Width: 6.2 inches  |  Height: Scaled to Frame\n(Paste or Insert Screenshot of Cognivista Application Here)")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(10)
    r2.italic = True
    r2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(28)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(6)
    p_cap.paragraph_format.space_after = Pt(10)
    rc = p_cap.add_run(figure_caption)
    rc.bold = True
    rc.font.name = 'Times New Roman'
    rc.font.size = Pt(11)
    rc.font.color.rgb = RGBColor(0, 0, 0)

print("Helper functions compiled successfully.")
